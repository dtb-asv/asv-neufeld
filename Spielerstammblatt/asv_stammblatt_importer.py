import os, re, shutil, tempfile
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import pymupdf
    import pytesseract
    from PIL import Image, ImageTk, ImageOps
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError as e:
    raise SystemExit(f"Fehlendes Python-Paket: {e}. Bitte 'pip install pymupdf pytesseract pillow openpyxl' ausführen.")

FIELDS = [
    ("mannschaft", "Mannschaft"), ("vorname", "Vorname"), ("nachname", "Nachname"),
    ("geburtsdatum", "Geburtsdatum"), ("adresse", "Adresse"), ("plz", "PLZ"), ("ort", "Ort"),
    ("erz1", "Erziehungsberechtigte/r 1"), ("email1", "E-Mail 1"), ("telefon1", "Telefon 1"),
    ("erz2", "Erziehungsberechtigte/r 2"), ("email2", "E-Mail 2"), ("telefon2", "Telefon 2"),
    ("unterschrift_ort", "Unterschrift Ort"), ("unterschrift_datum", "Unterschrift Datum"),
]
HEADERS = [label.upper().replace("/", "_") for _, label in FIELDS] + ["QUELLDATEI", "PDF_SEITE"]

TEAMS = ["U6", "U7", "U8", "U9", "U10", "U11", "U12", "U13", "U14", "U15", "U16", "U17", "U18", "U23", "KM"]


def setup_tesseract():
    if shutil.which("tesseract"):
        return True
    candidates = [
        r"C:\\Program Files\\Tesseract-OCR\\tesseract.exe",
        r"C:\\Program Files (x86)\\Tesseract-OCR\\tesseract.exe",
    ]
    for p in candidates:
        if os.path.exists(p):
            pytesseract.pytesseract.tesseract_cmd = p
            return True
    return False


def ocr_image(img):
    if not setup_tesseract():
        raise RuntimeError("Tesseract OCR wurde nicht gefunden. Siehe README_ASV_Stammblatt.txt")
    gray = ImageOps.grayscale(img)
    gray = ImageOps.autocontrast(gray)
    try:
        return pytesseract.image_to_string(gray, lang="deu+eng", config="--psm 6")
    except Exception:
        return pytesseract.image_to_string(gray, lang="eng", config="--psm 6")


def clean(v):
    v = (v).replace("ﬁ", "fi").replace("ﬂ", "fl")
    v = re.sub(r"_+", " ", v)
    return re.sub(r"\s+", " ", v).strip(" :-_\t")


def next_value(lines, label_index, stop_words=()):
    for j in range(label_index + 1, min(len(lines), label_index + 5)):
        s = clean(lines[j])
        if not s:
            continue
        low = s.lower()
        if any(w in low for w in stop_words):
            return ""
        # skip template label remnants
        if low.startswith(("vorname", "nachname", "geburtsdatum", "adresse", "plz / ort", "name", "e-mail", "telefon")):
            continue
        return s
    return ""


def parse_text(text):
    # Parse the known ASV form. Everything remains editable in the GUI.
    raw = [x.strip() for x in text.replace("\r", "").split("\n")]
    lines = [x for x in raw if x.strip()]
    d = {k: "" for k, _ in FIELDS}

    # Digitally filled version: PyMuPDF often extracts the typed values as one block
    # after the printed form. This is much more reliable than OCR for those PDFs.
    sig_line = -1
    for ix, ln in enumerate(lines):
        if ln.lower().startswith("unterschrift erziehungsberechtigte/r"):
            sig_line = ix
            break
    if sig_line >= 0:
        tail = [clean(x) for x in lines[sig_line+1:] if clean(x)]
        if len(tail) >= 11:
            # Expected order of filled fields in the ASV template.
            vals = tail[:13]
            keys = ["vorname","nachname","geburtsdatum","adresse","plzort","erz1","email1","telefon1","erz2","email2","telefon2","unterschrift_datum","unterschrift_ort"]
            temp = dict(zip(keys, vals))
            for k in ("vorname","nachname","geburtsdatum","adresse","erz1","email1","telefon1","erz2","email2","telefon2","unterschrift_datum","unterschrift_ort"):
                d[k] = temp.get(k, "")
            m = re.match(r"^(\d{4})\s+(.+)$", temp.get("plzort", ""))
            if m:
                d["plz"], d["ort"] = m.group(1), m.group(2)
            return d

    def find_idx(pattern, start=0):
        rg = re.compile(pattern, re.I)
        for i in range(start, len(lines)):
            if rg.search(lines[i]): return i
        return -1

    i = find_idx(r"^Vorname\s*:")
    if i >= 0: d["vorname"] = next_value(lines, i)
    i = find_idx(r"^Nachname\s*:")
    if i >= 0: d["nachname"] = next_value(lines, i)
    i = find_idx(r"^Geburtsdatum\s*:")
    if i >= 0: d["geburtsdatum"] = next_value(lines, i)
    i = find_idx(r"^Adresse\s*:")
    if i >= 0:
        d["adresse"] = next_value(lines, i, ("plz",))
        # Digital form often places "2491 Neufeld" on next line after address value, before PLZ label.
        if i + 2 < len(lines):
            cand = clean(lines[i+2])
            m = re.match(r"^(\d{4})\s+(.+)$", cand)
            if m:
                d["plz"], d["ort"] = m.group(1), m.group(2)
    i_plz = find_idx(r"^PLZ\s*/\s*Ort\s*:")
    if i_plz >= 0 and (not d["plz"] or not d["ort"]):
        val = next_value(lines, i_plz)
        m = re.match(r"^(\d{4})\s+(.+)$", val)
        if m: d["plz"], d["ort"] = m.group(1), m.group(2)

    sec1 = find_idx(r"Erziehungsberechtigte/r\s*1")
    sec2 = find_idx(r"Erziehungsberechtigte/r\s*2")
    sec3 = find_idx(r"EINVERSTÄNDNIS")

    def parse_guardian(start, end, suffix):
        if start < 0: return
        end = end if end >= 0 else len(lines)
        for field, pat in [("erz", r"^Name\s*:"), ("email", r"^E-Mail\s*:"), ("telefon", r"^Telefon\s*:")]:
            for idx in range(start+1, end):
                if re.search(pat, lines[idx], re.I):
                    d[field+suffix] = next_value(lines, idx)
                    break
    parse_guardian(sec1, sec2, "1")
    parse_guardian(sec2, sec3, "2")

    # Signature location/date: digital PDFs may put values on following lines in reverse visual extraction order.
    sig_start = sec3 if sec3 >= 0 else 0
    date_candidates = re.findall(r"\b\d{1,2}[./-]\d{1,2}[./-]\d{2,4}\b", "\n".join(lines[sig_start:]))
    if date_candidates: d["unterschrift_datum"] = date_candidates[-1]
    oi = find_idx(r"^Ort\s*:", sig_start)
    if oi >= 0:
        candidates = []
        for s in lines[oi+1:oi+6]:
            c = clean(s)
            if not c or re.search(r"Datum|Unterschrift|\d{1,2}[./-]\d{1,2}", c, re.I): continue
            candidates.append(c)
        if candidates: d["unterschrift_ort"] = candidates[-1]
    return d


def ocr_field(img, box, kind="text"):
    """OCR only the handwritten answer area of a fixed ASV form.
    Coordinates are relative fractions (left, top, right, bottom).
    """
    w, h = img.size
    crop = img.crop((int(box[0]*w), int(box[1]*h), int(box[2]*w), int(box[3]*h)))
    crop = ImageOps.grayscale(crop)
    crop = ImageOps.autocontrast(crop)
    # enlarge strongly; field crops contain far less printed template text
    crop = crop.resize((crop.width*3, crop.height*3), Image.Resampling.LANCZOS)
    config = "--psm 7"
    if kind == "digits":
        config += " -c tessedit_char_whitelist=0123456789./-"
    elif kind == "phone":
        config += " -c tessedit_char_whitelist=0123456789+()/-"
    try:
        txt = pytesseract.image_to_string(crop, lang="deu+eng", config=config)
    except Exception:
        txt = pytesseract.image_to_string(crop, lang="eng", config=config)
    return clean(txt)


def parse_fixed_scan(img):
    """Read the current ASV Neufeld form field-by-field instead of OCRing the whole page.
    This deliberately avoids labels/section headings that confused V2.
    """
    # Coordinates tuned to the uploaded current ASV form. They include only the answer line.
    boxes = {
        "vorname": (0.235, 0.278, 0.570, 0.307),
        "nachname": (0.235, 0.307, 0.570, 0.336),
        "geburtsdatum": (0.235, 0.336, 0.570, 0.365),
        "adresse": (0.235, 0.365, 0.570, 0.394),
        "plzort": (0.235, 0.394, 0.570, 0.423),
        "erz1": (0.235, 0.500, 0.570, 0.530),
        "email1": (0.235, 0.530, 0.570, 0.559),
        "telefon1": (0.235, 0.559, 0.570, 0.588),
        "erz2": (0.235, 0.616, 0.570, 0.646),
        "email2": (0.235, 0.646, 0.570, 0.675),
        "telefon2": (0.235, 0.675, 0.570, 0.704),
        "unterschrift_ort": (0.235, 0.866, 0.390, 0.895),
        "unterschrift_datum": (0.455, 0.866, 0.625, 0.895),
    }
    d = {k: "" for k, _ in FIELDS}
    for k, box in boxes.items():
        kind = "text"
        if k in ("geburtsdatum", "unterschrift_datum"): kind = "digits"
        if k.startswith("telefon"): kind = "phone"
        d[k] = ocr_field(img, box, kind)
    plzort = d.pop("plzort", "") if "plzort" in d else ocr_field(img, boxes["plzort"])
    m = re.search(r"\b(\d{4})\b\s*(.*)", plzort)
    if m:
        d["plz"] = m.group(1)
        d["ort"] = clean(m.group(2))
    else:
        # Keep uncertain OCR visible for manual correction instead of losing it.
        d["ort"] = plzort
    return d


def pdf_pages(path):
    """Liest alle Seiten einer PDF und übernimmt nur echte Spielerstammblätter.

    Bei Scan-PDFs wird zuerst nur der Kopfbereich per OCR geprüft. Dadurch werden
    angehängte DSGVO-Seiten schnell erkannt und gar nicht als Spieler importiert.
    Erst bei einer Stammblatt-Seite wird die ganze Seite per OCR gelesen.
    """
    doc = pymupdf.open(path)
    out = []
    for n, page in enumerate(doc):
        direct = page.get_text("text") or ""
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2.5,2.5), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

        direct_upper = direct.upper()
        if "SPIELERSTAMMBLATT" in direct_upper:
            is_form = True
            text = direct
        elif "DATENSCHUTZHINWEISE" in direct_upper or "DSGVO" in direct_upper[:500]:
            is_form = False
            text = direct
        else:
            # Scan: nur den oberen Bereich prüfen. Dort steht beim Formular
            # deutlich SPIELERSTAMMBLATT, bei der Folgeseite Datenschutzhinweise/DSGVO.
            head = img.crop((0, 0, img.width, int(img.height * 0.28)))
            head_text = ocr_image(head)
            hu = head_text.upper()
            is_form = ("SPIELERSTAMMBLATT" in hu or
                       ("ANGABEN" in hu and "SPIELER" in hu))
            text = ocr_image(img) if is_form else head_text

        if is_form:
            # Digital PDF: parse embedded text. Scan: fixed-position field OCR.
            if "SPIELERSTAMMBLATT" in direct_upper:
                data = parse_text(direct)
            else:
                data = parse_fixed_scan(img)
            out.append((n+1, data, img.copy()))
    doc.close()
    return out

def image_page(path):
    img = Image.open(path).convert("RGB")
    return [(1, parse_fixed_scan(img), img.copy())]


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ASV Neufeld – Spielerstammblatt Import V3")
        self.geometry("1280x820")
        self.minsize(1050, 700)
        self.records = []
        self.idx = -1
        self.photo = None
        self.vars = {k: tk.StringVar() for k, _ in FIELDS}
        self.status = tk.StringVar(value="Bereit")
        self.build()

    def build(self):
        top = ttk.Frame(self, padding=10); top.pack(fill="x")
        ttk.Label(top, text="Mannschaft:", font=("Segoe UI", 10, "bold")).pack(side="left")
        self.team = ttk.Combobox(top, values=TEAMS, width=10, state="normal")
        self.team.set("U10"); self.team.pack(side="left", padx=(6,18))
        ttk.Button(top, text="Stammblätter auswählen", command=self.choose_files).pack(side="left", padx=4)
        ttk.Button(top, text="Ordner auswählen", command=self.choose_folder).pack(side="left", padx=4)
        ttk.Button(top, text="Excel exportieren", command=self.export_excel).pack(side="right", padx=4)

        nav = ttk.Frame(self, padding=(10,0,10,8)); nav.pack(fill="x")
        self.counter = ttk.Label(nav, text="Noch keine Stammblätter geladen")
        self.counter.pack(side="left")
        ttk.Button(nav, text="◀ Vorheriger", command=lambda:self.move(-1)).pack(side="right", padx=3)
        ttk.Button(nav, text="Nächster ▶", command=lambda:self.move(1)).pack(side="right", padx=3)

        pan = ttk.Panedwindow(self, orient="horizontal"); pan.pack(fill="both", expand=True, padx=10, pady=5)
        left = ttk.Frame(pan); right = ttk.Frame(pan); pan.add(left, weight=3); pan.add(right, weight=2)
        self.canvas = tk.Canvas(left, bg="#d9d9d9", highlightthickness=0)
        self.canvas.pack(fill="both", expand=True)
        self.canvas.bind("<Configure>", lambda e:self.show_image())

        form = ttk.Frame(right, padding=12); form.pack(fill="both", expand=True)
        ttk.Label(form, text="Erkannte Daten – bitte kontrollieren", font=("Segoe UI", 14, "bold")).grid(row=0,column=0,columnspan=2,sticky="w",pady=(0,12))
        for r,(key,label) in enumerate(FIELDS, start=1):
            ttk.Label(form, text=label+":").grid(row=r,column=0,sticky="w",pady=3,padx=(0,8))
            if key == "mannschaft":
                w = ttk.Combobox(form, textvariable=self.vars[key], values=TEAMS, state="normal")
            else:
                w = ttk.Entry(form, textvariable=self.vars[key])
            w.grid(row=r,column=1,sticky="ew",pady=3)
        form.columnconfigure(1, weight=1)
        ttk.Label(form, text="Hinweis: Handschriftliche OCR kann Fehler enthalten.\nVor dem Export bitte die Felder mit dem Original links vergleichen.", foreground="#8a4b00").grid(row=len(FIELDS)+1,column=0,columnspan=2,sticky="w",pady=(15,0))

        bottom = ttk.Frame(self, padding=10); bottom.pack(fill="x")
        ttk.Label(bottom, textvariable=self.status).pack(side="left")

    def choose_files(self):
        paths = filedialog.askopenfilenames(title="Stammblätter auswählen", filetypes=[("Stammblätter", "*.pdf *.jpg *.jpeg *.png"), ("PDF", "*.pdf"), ("Bilder", "*.jpg *.jpeg *.png")])
        if paths: self.process(paths)

    def choose_folder(self):
        folder = filedialog.askdirectory(title="Ordner mit Stammblättern auswählen")
        if folder:
            paths = [str(p) for p in sorted(Path(folder).iterdir()) if p.suffix.lower() in (".pdf",".jpg",".jpeg",".png")]
            if not paths: messagebox.showinfo("Keine Dateien", "In diesem Ordner wurden keine PDF/JPG/PNG-Dateien gefunden."); return
            self.process(paths)

    def process(self, paths):
        self.save_current()
        team = self.team.get().strip()
        added = 0
        self.config(cursor="watch"); self.update_idletasks()
        try:
            for pos,path in enumerate(paths,1):
                self.status.set(f"Verarbeite {pos}/{len(paths)}: {Path(path).name}"); self.update()
                try:
                    pages = pdf_pages(path) if Path(path).suffix.lower()==".pdf" else image_page(path)
                    for page_no,d,img in pages:
                        d["mannschaft"] = team
                        self.records.append({"data":d, "image":img, "source":Path(path).name, "page":page_no})
                        added += 1
                except Exception as e:
                    messagebox.showwarning("Datei übersprungen", f"{Path(path).name}\n\n{e}")
        finally:
            self.config(cursor="")
        if added:
            self.idx = len(self.records)-added
            self.load_current()
            self.status.set(f"{added} Stammblatt/Stammblätter geladen. Bitte Daten kontrollieren.")
        else: self.status.set("Keine Stammblätter geladen.")

    def save_current(self):
        if 0 <= self.idx < len(self.records):
            self.records[self.idx]["data"] = {k:v.get().strip() for k,v in self.vars.items()}

    def load_current(self):
        if not (0 <= self.idx < len(self.records)): return
        d=self.records[self.idx]["data"]
        for k in self.vars: self.vars[k].set(d.get(k,""))
        r=self.records[self.idx]
        self.counter.config(text=f"{self.idx+1} von {len(self.records)}  –  {r['source']}  (Seite {r['page']})")
        self.show_image()

    def move(self, delta):
        if not self.records: return
        self.save_current(); self.idx=max(0,min(len(self.records)-1,self.idx+delta)); self.load_current()

    def show_image(self):
        if not (0 <= self.idx < len(self.records)): return
        img=self.records[self.idx]["image"].copy()
        cw=max(self.canvas.winfo_width()-20,100); ch=max(self.canvas.winfo_height()-20,100)
        img.thumbnail((cw,ch), Image.Resampling.LANCZOS)
        self.photo=ImageTk.PhotoImage(img)
        self.canvas.delete("all"); self.canvas.create_image(cw//2+10,ch//2+10,image=self.photo,anchor="center")

    def export_excel(self):
        if not self.records: messagebox.showinfo("Keine Daten", "Bitte zuerst Stammblätter laden."); return
        self.save_current()
        path=filedialog.asksaveasfilename(title="Excel speichern", defaultextension=".xlsx", initialfile="ASV_Spielerstamm.xlsx", filetypes=[("Excel", "*.xlsx")])
        if not path: return
        if os.path.exists(path):
            wb=load_workbook(path); ws=wb.active
            if ws.max_row==1 and ws.cell(1,1).value is None: ws.append(HEADERS)
        else:
            wb=Workbook(); ws=wb.active; ws.title="Spieler"; ws.append(HEADERS)
        # Header styling
        for c in ws[1]:
            c.font=Font(bold=True, color="FFFFFF"); c.fill=PatternFill("solid", fgColor="1F4E78"); c.alignment=Alignment(horizontal="center")
        for rec in self.records:
            d=rec["data"]
            ws.append([d.get(k,"") for k,_ in FIELDS]+[rec["source"], rec["page"]])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        widths=[14,18,22,15,28,9,20,28,30,20,28,30,20,18,18,32,12]
        for i,w in enumerate(widths,1): ws.column_dimensions[chr(64+i) if i<=26 else "A"].width=w
        wb.save(path)
        self.status.set(f"Excel gespeichert: {path}")
        messagebox.showinfo("Fertig", f"{len(self.records)} Datensätze wurden nach Excel exportiert.")

if __name__ == "__main__":
    App().mainloop()
