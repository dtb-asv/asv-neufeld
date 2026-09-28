import os, re
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

try:
    import pymupdf
    from PIL import Image, ImageTk
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
except ImportError as e:
    raise SystemExit(f"Fehlendes Python-Paket: {e}. Bitte INSTALL_PYTHON_PAKETE.bat ausführen.")

FIELDS = [
    ("mannschaft", "Mannschaft"), ("vorname", "Vorname"), ("nachname", "Nachname"),
    ("geburtsdatum", "Geburtsdatum"), ("adresse", "Adresse"), ("plz", "PLZ"), ("ort", "Ort"),
    ("erz1", "Erziehungsberechtigte/r 1"), ("email1", "E-Mail 1"), ("telefon1", "Telefon 1"),
    ("erz2", "Erziehungsberechtigte/r 2"), ("email2", "E-Mail 2"), ("telefon2", "Telefon 2"),
    ("unterschrift_ort", "Unterschrift Ort"), ("unterschrift_datum", "Unterschrift Datum"),
]
HEADERS = [label.upper().replace("/", "_") for _, label in FIELDS] + ["QUELLDATEI"]
TEAMS = ["U6", "U7", "U8", "U9", "U10", "U11", "U12", "U13", "U14", "U15", "U16", "U17", "U18", "U23", "KM"]


def clean(v):
    return re.sub(r"\s+", " ", (v or "").replace("ﬁ", "fi").replace("ﬂ", "fl")).strip(" :-_\t")


def parse_digital_asv(text):
    """Conservative parser for the known digitally-filled ASV form. Returns blank fields if unsure."""
    d = {k: "" for k, _ in FIELDS}
    lines = [clean(x) for x in (text or "").replace("\r", "").split("\n") if clean(x)]
    sig = next((i for i,x in enumerate(lines) if x.lower().startswith("unterschrift erziehungsberechtigte/r")), -1)
    if sig < 0:
        return d
    tail = [x for x in lines[sig+1:] if x]
    # Known editable PDF exports entered values after the printed form in this order.
    if len(tail) < 11:
        return d
    vals = tail[:13]
    keys = ["vorname","nachname","geburtsdatum","adresse","plzort","erz1","email1","telefon1","erz2","email2","telefon2","unterschrift_datum","unterschrift_ort"]
    t = dict(zip(keys, vals))
    for k in ("vorname","nachname","geburtsdatum","adresse","erz1","email1","telefon1","erz2","email2","telefon2","unterschrift_datum","unterschrift_ort"):
        d[k] = t.get(k, "")
    m = re.match(r"^(\d{4})\s+(.+)$", t.get("plzort", ""))
    if m:
        d["plz"], d["ort"] = m.group(1), m.group(2)
    return d


def pdf_form_pages(path):
    """Load likely form pages. Digital appended DSGVO pages are skipped; scans are kept for manual entry."""
    doc = pymupdf.open(path)
    out = []
    for n, page in enumerate(doc):
        text = page.get_text("text") or ""
        compact = re.sub(r"\s", "", text)
        digital = len(compact) >= 80
        if digital:
            upper = text.upper()
            is_form = "SPIELERSTAMMBLATT" in upper
            if not is_form:
                continue
        else:
            # No OCR by design. For image-only PDFs, first page is assumed to be the submitted form.
            # Extra image-only pages can be removed with "Stammblatt entfernen" if needed.
            if n > 0:
                continue
        pix = page.get_pixmap(matrix=pymupdf.Matrix(2.2, 2.2), alpha=False)
        img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
        data = parse_digital_asv(text) if digital else {k:"" for k,_ in FIELDS}
        out.append((n+1, img, data, digital))
    doc.close()
    return out


def image_form(path):
    img = Image.open(path).convert("RGB")
    return [(1, img.copy(), {k:"" for k,_ in FIELDS}, False)]


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("ASV Neufeld – Spielerstammblatt Erfassung V4")
        self.geometry("1380x860")
        self.minsize(1100, 720)
        self.records, self.idx, self.photo = [], -1, None
        self.zoom = 1.0
        self.vars = {k: tk.StringVar() for k,_ in FIELDS}
        self.status = tk.StringVar(value="Bereit – keine OCR, kein API-Key. Digitale PDFs werden soweit möglich vorausgefüllt.")
        self.entries = []
        self.build()

    def build(self):
        top = ttk.Frame(self, padding=10); top.pack(fill="x")
        ttk.Label(top, text="Mannschaft für neue Dateien:", font=("Segoe UI",10,"bold")).pack(side="left")
        self.team = ttk.Combobox(top, values=TEAMS, width=10, state="normal"); self.team.set("U10"); self.team.pack(side="left", padx=(6,18))
        ttk.Button(top, text="Dateien auswählen", command=self.choose_files).pack(side="left", padx=4)
        ttk.Button(top, text="Ordner auswählen", command=self.choose_folder).pack(side="left", padx=4)
        ttk.Button(top, text="Excel exportieren", command=self.export_excel).pack(side="right", padx=4)

        nav = ttk.Frame(self, padding=(10,0,10,8)); nav.pack(fill="x")
        self.counter = ttk.Label(nav, text="Noch keine Stammblätter geladen"); self.counter.pack(side="left")
        ttk.Button(nav, text="−", width=3, command=lambda:self.change_zoom(-0.15)).pack(side="left", padx=(20,2))
        ttk.Button(nav, text="100%", command=self.reset_zoom).pack(side="left", padx=2)
        ttk.Button(nav, text="+", width=3, command=lambda:self.change_zoom(0.15)).pack(side="left", padx=2)
        ttk.Button(nav, text="Stammblatt entfernen", command=self.remove_current).pack(side="right", padx=(12,3))
        ttk.Button(nav, text="◀ Vorheriger", command=lambda:self.move(-1)).pack(side="right", padx=3)
        ttk.Button(nav, text="Speichern && Nächster ▶", command=self.save_and_next).pack(side="right", padx=3)

        pan = ttk.Panedwindow(self, orient="horizontal"); pan.pack(fill="both", expand=True, padx=10, pady=5)
        left = ttk.Frame(pan); right = ttk.Frame(pan); pan.add(left, weight=3); pan.add(right, weight=2)
        self.canvas = tk.Canvas(left, bg="#d7d7d7", highlightthickness=0)
        ybar=ttk.Scrollbar(left, orient="vertical", command=self.canvas.yview); xbar=ttk.Scrollbar(left, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=ybar.set, xscrollcommand=xbar.set)
        self.canvas.grid(row=0,column=0,sticky="nsew"); ybar.grid(row=0,column=1,sticky="ns"); xbar.grid(row=1,column=0,sticky="ew")
        left.rowconfigure(0,weight=1); left.columnconfigure(0,weight=1)
        self.canvas.bind("<Configure>", lambda e:self.show_image())

        form = ttk.Frame(right, padding=14); form.pack(fill="both", expand=True)
        ttk.Label(form, text="Daten erfassen / kontrollieren", font=("Segoe UI",15,"bold")).grid(row=0,column=0,columnspan=2,sticky="w",pady=(0,4))
        self.mode_label=ttk.Label(form, text="", foreground="#555555"); self.mode_label.grid(row=1,column=0,columnspan=2,sticky="w",pady=(0,12))
        for r,(key,label) in enumerate(FIELDS, start=2):
            ttk.Label(form,text=label+":").grid(row=r,column=0,sticky="w",pady=4,padx=(0,10))
            if key=="mannschaft": w=ttk.Combobox(form,textvariable=self.vars[key],values=TEAMS,state="normal")
            else: w=ttk.Entry(form,textvariable=self.vars[key])
            w.grid(row=r,column=1,sticky="ew",pady=4); self.entries.append(w)
        form.columnconfigure(1,weight=1)
        ttk.Label(form,text="Tipp: Mit TAB von Feld zu Feld. Nach dem letzten Feld: Alt+N für Speichern & Nächster.",foreground="#555555").grid(row=len(FIELDS)+2,column=0,columnspan=2,sticky="w",pady=(16,0))
        self.bind_all("<Alt-n>", lambda e:self.save_and_next())
        self.bind_all("<Alt-N>", lambda e:self.save_and_next())

        bottom=ttk.Frame(self,padding=10); bottom.pack(fill="x"); ttk.Label(bottom,textvariable=self.status).pack(side="left")

    def choose_files(self):
        paths=filedialog.askopenfilenames(title="Stammblätter auswählen",filetypes=[("Stammblätter","*.pdf *.jpg *.jpeg *.png"),("PDF","*.pdf"),("Bilder","*.jpg *.jpeg *.png")])
        if paths:self.process(paths)

    def choose_folder(self):
        folder=filedialog.askdirectory(title="Ordner mit Stammblättern auswählen")
        if folder:
            paths=[str(p) for p in sorted(Path(folder).iterdir()) if p.suffix.lower() in (".pdf",".jpg",".jpeg",".png")]
            if not paths: messagebox.showinfo("Keine Dateien","Keine PDF/JPG/PNG-Dateien gefunden."); return
            self.process(paths)

    def process(self,paths):
        self.save_current(); team=self.team.get().strip(); start=len(self.records); skipped=[]
        self.config(cursor="watch")
        try:
            for pos,path in enumerate(paths,1):
                self.status.set(f"Lade {pos}/{len(paths)}: {Path(path).name}"); self.update()
                try:
                    pages=pdf_form_pages(path) if Path(path).suffix.lower()==".pdf" else image_form(path)
                    for page_no,img,data,digital in pages:
                        data["mannschaft"]=team
                        self.records.append({"data":data,"image":img,"source":Path(path).name,"page":page_no,"digital":digital})
                except Exception as e: skipped.append(f"{Path(path).name}: {e}")
        finally:self.config(cursor="")
        added=len(self.records)-start
        if added:
            self.idx=start; self.load_current(); self.status.set(f"{added} Stammblatt/Stammblätter geladen.")
        if skipped: messagebox.showwarning("Einige Dateien konnten nicht geladen werden","\n".join(skipped[:8]))

    def save_current(self):
        if 0<=self.idx<len(self.records): self.records[self.idx]["data"]={k:v.get().strip() for k,v in self.vars.items()}

    def load_current(self):
        if not (0<=self.idx<len(self.records)): return
        r=self.records[self.idx]; d=r["data"]
        for k in self.vars:self.vars[k].set(d.get(k,""))
        self.counter.config(text=f"{self.idx+1} von {len(self.records)} – {r['source']} (Seite {r['page']})")
        self.mode_label.config(text="Digitales PDF: erkannte Werte bitte kontrollieren." if r["digital"] else "Scan/Bild: bitte Daten rechts eingeben.")
        self.zoom=1.0; self.show_image()
        if self.entries: self.entries[1 if len(self.entries)>1 else 0].focus_set()

    def move(self,delta):
        if not self.records:return
        self.save_current(); self.idx=max(0,min(len(self.records)-1,self.idx+delta)); self.load_current()

    def save_and_next(self):
        if not self.records:return
        self.save_current()
        if self.idx < len(self.records)-1:
            self.idx+=1; self.load_current(); self.status.set("Gespeichert – nächstes Stammblatt.")
        else:self.status.set("Letztes Stammblatt gespeichert. Jetzt Excel exportieren.")

    def remove_current(self):
        if not self.records:return
        if not messagebox.askyesno("Entfernen","Dieses Stammblatt aus der aktuellen Erfassung entfernen?"):return
        self.records.pop(self.idx)
        if not self.records:
            self.idx=-1; self.canvas.delete("all"); self.counter.config(text="Noch keine Stammblätter geladen")
            for v in self.vars.values():v.set("")
            return
        self.idx=min(self.idx,len(self.records)-1); self.load_current()

    def change_zoom(self,delta):
        self.zoom=max(0.5,min(2.5,self.zoom+delta)); self.show_image()
    def reset_zoom(self): self.zoom=1.0; self.show_image()

    def show_image(self):
        if not (0<=self.idx<len(self.records)):return
        src=self.records[self.idx]["image"]
        cw=max(self.canvas.winfo_width()-30,200); ch=max(self.canvas.winfo_height()-30,200)
        base=min(cw/src.width,ch/src.height)
        scale=base*self.zoom
        size=(max(1,int(src.width*scale)),max(1,int(src.height*scale)))
        img=src.resize(size,Image.Resampling.LANCZOS)
        self.photo=ImageTk.PhotoImage(img); self.canvas.delete("all"); self.canvas.create_image(10,10,image=self.photo,anchor="nw")
        self.canvas.configure(scrollregion=(0,0,size[0]+20,size[1]+20))

    def export_excel(self):
        if not self.records: messagebox.showinfo("Keine Daten","Bitte zuerst Stammblätter laden."); return
        self.save_current()
        path=filedialog.asksaveasfilename(title="Excel speichern",defaultextension=".xlsx",initialfile="ASV_Spielerstamm.xlsx",filetypes=[("Excel","*.xlsx")])
        if not path:return
        if os.path.exists(path):
            wb=load_workbook(path); ws=wb.active
            if ws.max_row==1 and ws.cell(1,1).value is None:ws.append(HEADERS)
        else:
            wb=Workbook(); ws=wb.active; ws.title="Spieler"; ws.append(HEADERS)
        for c in ws[1]:
            c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78"); c.alignment=Alignment(horizontal="center")
        for rec in self.records:
            d=rec["data"]; ws.append([d.get(k,"") for k,_ in FIELDS]+[rec["source"]])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        widths=[14,18,22,15,28,9,20,28,30,20,28,30,20,18,18,32]
        for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
        wb.save(path); self.status.set(f"Excel gespeichert: {path}"); messagebox.showinfo("Fertig",f"{len(self.records)} Datensätze wurden nach Excel exportiert.")

if __name__=="__main__": App().mainloop()
